"""File-backed COC template lookup and XLSX export implementation."""

import os
import posixpath
import re
import shutil
import tempfile
from copy import deepcopy
from datetime import date, datetime
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZIP_DEFLATED, ZipFile

from config.app_config import (
    COC_TEMPLATE_45_FILENAME,
    COC_TEMPLATE_48_FILENAME,
    COC_TEMPLATE_FILENAME,
    DEFAULT_PART_LOOKUP_ROOT,
)

_INVALID_FILENAME_CHARACTERS = re.compile(r'[<>:"/\\|?*]')


def normalise_serial(value: str) -> str:
    """Return a display/search serial without an optional SN prefix."""
    serial = str(value or "").strip()
    if serial[:2].upper() == "SN" and serial[2:]:
        return serial[2:]
    return serial


def find_part_number(main_board_serial, lookup_root=DEFAULT_PART_LOOKUP_ROOT):
    """Find the part number in the matching SN<serial>_<part> directory."""
    serial = normalise_serial(main_board_serial)
    if not serial:
        raise ValueError("Enter a Main Board serial number first.")

    root = Path(lookup_root)
    if not root.is_dir():
        raise FileNotFoundError(
            "The part-number lookup folder is not available:\n%s" % root
        )

    prefix = "SN%s_" % serial
    matches = sorted(
        directory
        for directory in root.iterdir()
        if directory.is_dir()
        and directory.name.casefold().startswith(prefix.casefold())
    )
    if not matches:
        raise LookupError(
            "No OSX-150 folder was found for Main Board serial %s." % serial
        )
    if len(matches) > 1:
        raise LookupError(
            "More than one OSX-150 folder was found for Main Board serial %s:\n%s"
            % (serial, "\n".join(directory.name for directory in matches))
        )

    part_number = matches[0].name[len(prefix):]
    if not part_number:
        raise LookupError("The matching folder does not contain a part number.")
    return part_number


def _safe_filename_component(value):
    component = _INVALID_FILENAME_CHARACTERS.sub("-", str(value).strip())
    component = component.rstrip(" .")
    return component or "Unknown"


def _next_output_path(directory, serial, export_timestamp=None):
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    timestamp = (export_timestamp or datetime.now()).strftime("%y%m%d-%H%M%S")
    serial_component = _safe_filename_component(serial)
    base_name = "COC OSX-150 %s_%s.xlsx" % (serial_component, timestamp)
    output_path = directory / base_name
    suffix = 2
    while output_path.exists():
        output_path = directory / (
            "COC OSX-150 %s_%s (%d).xlsx"
            % (serial_component, timestamp, suffix)
        )
        suffix += 1
    return output_path


def measurement_rows(template_capacity=48):
    """Return the expected logical-channel rows for a template capacity."""
    if int(template_capacity) == 45:
        return list(range(11, 56))
    if int(template_capacity) == 48:
        return list(range(11, 56)) + list(range(67, 70))
    raise ValueError("COC template capacity must be 45 or 48 channels.")


def template_filename_for_channel_count(channel_count):
    """Return the approved OSX-150 template name for a front-panel count."""
    count = int(channel_count)
    if not 1 <= count <= 48:
        raise ValueError("Front-panel channel count must be between 1 and 48.")
    return COC_TEMPLATE_45_FILENAME if count <= 45 else COC_TEMPLATE_48_FILENAME


def inspect_coc_template_capacity(template_path):
    """Validate the expected sheet/channel map and return capacity 45 or 48."""
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError(
            "The XLSX COC exporter requires the openpyxl package."
        ) from error

    path = Path(template_path)
    if not path.is_file():
        raise FileNotFoundError("The COC template was not found:\n%s" % path)
    workbook = load_workbook(path, read_only=False, data_only=False)
    try:
        if "OSX Template" not in workbook.sheetnames:
            raise ValueError("The COC template does not contain an 'OSX Template' sheet.")
        sheet = workbook["OSX Template"]
        for channel, row in enumerate(range(11, 56), start=1):
            if sheet["B%d" % row].value != channel:
                raise ValueError(
                    "The COC template has an invalid channel map at channel %d."
                    % channel
                )
        second_section = [sheet["B%d" % row].value for row in range(67, 70)]
        if second_section == [46, 47, 48]:
            return 48
        if all(value in (None, "") for value in second_section):
            return 45
        raise ValueError(
            "The COC template must support either channels 1-45 or channels 1-48."
        )
    finally:
        workbook.close()


def _trim_report_rows(sheet, channel_count, template_capacity):
    """Remove unused report rows while explicitly repairing merged ranges."""
    if template_capacity == 45:
        last_row = 10 + channel_count
    else:
        last_row = 66 + (channel_count - 45)

    for merged_range in list(sheet.merged_cells.ranges):
        if merged_range.min_row > last_row:
            sheet.unmerge_cells(str(merged_range))
            continue
        if merged_range.max_row > last_row:
            min_col = merged_range.min_col
            max_col = merged_range.max_col
            min_row = merged_range.min_row
            sheet.unmerge_cells(str(merged_range))
            if min_row < last_row or min_col < max_col:
                sheet.merge_cells(
                    start_row=min_row,
                    start_column=min_col,
                    end_row=last_row,
                    end_column=max_col,
                )

    if sheet.max_row > last_row:
        sheet.delete_rows(last_row + 1, sheet.max_row - last_row)
    sheet.print_area = "$B$1:$H$%d" % last_row


def _restore_template_drawings(template_path, output_path):
    """Restore template drawings that openpyxl cannot round-trip reliably."""
    relationships_namespace = "http://schemas.openxmlformats.org/package/2006/relationships"
    office_relationships_namespace = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    content_types_namespace = "http://schemas.openxmlformats.org/package/2006/content-types"
    drawing_relationship_type = (
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing"
    )

    with ZipFile(template_path, "r") as template_zip:
        template_names = set(template_zip.namelist())
        drawing_parts = {
            name
            for name in template_names
            if name.startswith("xl/drawings/") or name.startswith("xl/media/")
        }
        if not drawing_parts:
            return

        worksheet_drawings = []
        for relationship_name in template_names:
            if not relationship_name.startswith("xl/worksheets/_rels/"):
                continue
            if not relationship_name.endswith(".rels"):
                continue
            relationship_root = ElementTree.fromstring(
                template_zip.read(relationship_name)
            )
            for relationship in relationship_root:
                if relationship.get("Type") != drawing_relationship_type:
                    continue
                worksheet_filename = relationship_name.rsplit("/", 1)[1][:-5]
                worksheet_name = "xl/worksheets/%s" % worksheet_filename
                worksheet_drawings.append(
                    (worksheet_name, relationship_name, relationship.get("Target"))
                )

        with ZipFile(output_path, "r") as output_zip:
            output_names = set(output_zip.namelist())
            replacement_names = set(drawing_parts)
            replacement_names.add("[Content_Types].xml")
            replacement_names.update(
                relationship_name
                for _, relationship_name, _ in worksheet_drawings
            )
            replacement_names.update(
                worksheet_name for worksheet_name, _, _ in worksheet_drawings
            )

            output_content_types = ElementTree.fromstring(
                output_zip.read("[Content_Types].xml")
            )
            template_content_types = ElementTree.fromstring(
                template_zip.read("[Content_Types].xml")
            )
            existing_defaults = {
                child.get("Extension")
                for child in output_content_types
                if child.tag.endswith("Default")
            }
            existing_overrides = {
                child.get("PartName")
                for child in output_content_types
                if child.tag.endswith("Override")
            }
            for child in template_content_types:
                if child.tag.endswith("Default"):
                    if (
                        child.get("Extension") not in existing_defaults
                        and child.get("Extension") == "png"
                    ):
                        output_content_types.append(deepcopy(child))
                        existing_defaults.add(child.get("Extension"))
                elif child.tag.endswith("Override"):
                    part_name = child.get("PartName", "")
                    if (
                        part_name.startswith("/xl/drawings/")
                        and part_name not in existing_overrides
                    ):
                        output_content_types.append(deepcopy(child))
                        existing_overrides.add(part_name)
            ElementTree.register_namespace("", content_types_namespace)
            content_types_bytes = ElementTree.tostring(
                output_content_types,
                encoding="utf-8",
                xml_declaration=True,
            )

            modified_worksheets = {}
            modified_relationships = {}
            for worksheet_name, relationship_name, target in worksheet_drawings:
                worksheet_bytes = output_zip.read(worksheet_name)
                relationship_bytes = (
                    output_zip.read(relationship_name)
                    if relationship_name in output_names
                    else None
                )
                if relationship_bytes:
                    relationship_root = ElementTree.fromstring(relationship_bytes)
                else:
                    relationship_root = ElementTree.Element(
                        "{%s}Relationships" % relationships_namespace
                    )

                drawing_relationship = next(
                    (
                        relationship
                        for relationship in relationship_root
                        if relationship.get("Type") == drawing_relationship_type
                    ),
                    None,
                )
                if drawing_relationship is None:
                    used_ids = {
                        relationship.get("Id")
                        for relationship in relationship_root
                    }
                    relationship_id = "rId1"
                    index = 1
                    while relationship_id in used_ids:
                        index += 1
                        relationship_id = "rId%d" % index
                    drawing_relationship = ElementTree.SubElement(
                        relationship_root,
                        "{%s}Relationship" % relationships_namespace,
                        {
                            "Id": relationship_id,
                            "Type": drawing_relationship_type,
                            "Target": target,
                        },
                    )
                relationship_id = drawing_relationship.get("Id")
                worksheet_text = worksheet_bytes.decode("utf-8")
                if "<drawing " not in worksheet_text:
                    if "xmlns:r=" not in worksheet_text:
                        worksheet_text = worksheet_text.replace(
                            "<worksheet ",
                            '<worksheet xmlns:r="%s" ' % office_relationships_namespace,
                            1,
                        )
                    worksheet_text = worksheet_text.replace(
                        "</worksheet>",
                        '<drawing r:id="%s"/></worksheet>' % relationship_id,
                    )
                modified_worksheets[worksheet_name] = worksheet_text.encode("utf-8")
                ElementTree.register_namespace("", relationships_namespace)
                modified_relationships[relationship_name] = ElementTree.tostring(
                    relationship_root,
                    encoding="utf-8",
                    xml_declaration=True,
                )

            merge_tempfile = tempfile.NamedTemporaryFile(
                prefix=".coc-merge-",
                suffix=".xlsx",
                dir=str(Path(output_path).parent),
                delete=False,
            )
            merge_path = Path(merge_tempfile.name)
            merge_tempfile.close()
            try:
                with ZipFile(merge_path, "w", ZIP_DEFLATED) as merged_zip:
                    for item in output_zip.infolist():
                        if item.filename in replacement_names:
                            continue
                        merged_zip.writestr(item, output_zip.read(item.filename))
                    merged_zip.writestr("[Content_Types].xml", content_types_bytes)
                    for name in drawing_parts:
                        merged_zip.writestr(name, template_zip.read(name))
                    for name, value in modified_worksheets.items():
                        merged_zip.writestr(name, value)
                    for name, value in modified_relationships.items():
                        merged_zip.writestr(name, value)
                # Windows cannot replace the original XLSX while the source
                # ZipFile handle is still open.
                output_zip.close()
                os.replace(merge_path, output_path)
            finally:
                if merge_path.exists():
                    merge_path.unlink()


def export_coc(
    template_path,
    output_directory,
    measurements,
    part_number,
    main_board_serial,
    export_date=None,
    tested_by="",
    front_panel_channel_count=48,
    export_timestamp=None,
):
    """Create a COC XLSX copy while preserving non-output template cells."""
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError(
            "The XLSX COC exporter requires the openpyxl package."
        ) from error

    template = Path(template_path)
    if not template.is_file():
        raise FileNotFoundError("The COC template was not found:\n%s" % template)
    part_number = str(part_number or "").strip()
    serial = normalise_serial(main_board_serial)
    if not part_number:
        raise ValueError("Enter or look up a Part Number before writing the COC.")
    if not serial:
        raise ValueError("Enter a Main Board serial number before writing the COC.")
    channel_count = int(front_panel_channel_count)
    if not 1 <= channel_count <= 48:
        raise ValueError("Front-panel channel count must be between 1 and 48.")
    template_capacity = inspect_coc_template_capacity(template)
    required_capacity = 45 if channel_count <= 45 else 48
    if template_capacity < required_capacity:
        raise ValueError(
            "The selected COC template supports only %d channels; %d are required."
            % (template_capacity, channel_count)
        )

    output_path = _next_output_path(
        output_directory,
        serial,
        export_timestamp=export_timestamp,
    )
    temporary_file = tempfile.NamedTemporaryFile(
        prefix=".coc-",
        suffix=".xlsx",
        dir=str(output_path.parent),
        delete=False,
    )
    temporary_path = Path(temporary_file.name)
    temporary_file.close()

    workbook = None
    try:
        shutil.copy2(template, temporary_path)
        workbook = load_workbook(temporary_path)
        if "OSX Template" not in workbook.sheetnames:
            raise ValueError("The COC template does not contain an 'OSX Template' sheet.")
        sheet = workbook["OSX Template"]

        # These are the top-left cells of the merged template ranges.
        sheet["B4"] = part_number
        sheet["D4"] = serial
        sheet["F4"] = export_date or date.today()
        sheet["G4"] = str(tested_by or "").strip()

        records_by_channel = {
            record.channel: record
            for record in measurements
            if 1 <= record.channel <= channel_count
        }
        for channel, row in enumerate(measurement_rows(template_capacity), start=1):
            if channel > channel_count:
                break
            record = records_by_channel.get(channel)
            sheet["D%d" % row] = record.loss_1310 if record else None
            sheet["E%d" % row] = record.loss_1550 if record else None

        _trim_report_rows(sheet, channel_count, template_capacity)

        workbook.save(temporary_path)
        workbook.close()
        workbook = None
        _restore_template_drawings(template, temporary_path)
        os.replace(temporary_path, output_path)
    finally:
        if workbook is not None:
            workbook.close()
        if temporary_path.exists():
            temporary_path.unlink()

    return output_path


__all__ = [
    "COC_TEMPLATE_45_FILENAME",
    "COC_TEMPLATE_48_FILENAME",
    "COC_TEMPLATE_FILENAME",
    "DEFAULT_PART_LOOKUP_ROOT",
    "export_coc",
    "find_part_number",
    "inspect_coc_template_capacity",
    "measurement_rows",
    "normalise_serial",
    "template_filename_for_channel_count",
]
