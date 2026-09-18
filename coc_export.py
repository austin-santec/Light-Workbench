"""Write completed or partial ILM runs into the XLSX COC template."""

import os
import re
import shutil
import tempfile
from datetime import date
from pathlib import Path


DEFAULT_PART_LOOKUP_ROOT = Path(
    r"U:\Product Log\Units-COCs-Param Files\OSX-150"
)
COC_TEMPLATE_FILENAME = "OSX-100 Single Mode COC Template 1.xlsx"
_INVALID_FILENAME_CHARACTERS = re.compile(r'[<>:"/\\|?*]')


def normalise_serial(value: str) -> str:
    """Return a display/search serial without an optional SN prefix."""
    serial = str(value or "").strip()
    if serial[:2].upper() == "SN" and serial[2:]:
        return serial[2:]
    return serial


def find_part_number(main_board_serial, lookup_root=DEFAULT_PART_LOOKUP_ROOT):
    """Find the part number in the matching SN<serial>_<part> directory."""
    serial = normalise_serial(main_board_serial)
    if not serial:
        raise ValueError("Enter a Main Board serial number first.")

    root = Path(lookup_root)
    if not root.is_dir():
        raise FileNotFoundError(
            "The part-number lookup folder is not available:\n%s" % root
        )

    prefix = "SN%s_" % serial
    matches = sorted(
        directory
        for directory in root.iterdir()
        if directory.is_dir() and directory.name.casefold().startswith(prefix.casefold())
    )
    if not matches:
        raise LookupError(
            "No OSX-150 folder was found for Main Board serial %s." % serial
        )
    if len(matches) > 1:
        raise LookupError(
            "More than one OSX-150 folder was found for Main Board serial %s:\n%s"
            % (serial, "\n".join(directory.name for directory in matches))
        )

    part_number = matches[0].name[len(prefix):]
    if not part_number:
        raise LookupError("The matching folder does not contain a part number.")
    return part_number


def _safe_filename_component(value):
    component = _INVALID_FILENAME_CHARACTERS.sub("-", str(value).strip())
    component = component.rstrip(" .")
    return component or "Unknown"


def _next_output_path(directory, serial):
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    base_name = "COC OSX-150 %s.xlsx" % _safe_filename_component(serial)
    output_path = directory / base_name
    suffix = 2
    while output_path.exists():
        output_path = directory / (
            "COC OSX-150 %s (%d).xlsx"
            % (_safe_filename_component(serial), suffix)
        )
        suffix += 1
    return output_path


def _measurement_rows():
    """Return the COC rows for logical channels 1 through 48."""
    return list(range(11, 56)) + list(range(67, 70))


def export_coc(
    template_path,
    output_directory,
    measurements,
    part_number,
    main_board_serial,
    export_date=None,
):
    """Create a COC XLSX copy while preserving all non-output template cells."""
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError(
            "The XLSX COC exporter requires the openpyxl package."
        ) from error

    template = Path(template_path)
    if not template.is_file():
        raise FileNotFoundError("The COC template was not found:\n%s" % template)
    part_number = str(part_number or "").strip()
    serial = normalise_serial(main_board_serial)
    if not part_number:
        raise ValueError("Enter or look up a Part Number before writing the COC.")
    if not serial:
        raise ValueError("Enter a Main Board serial number before writing the COC.")

    output_path = _next_output_path(output_directory, serial)
    temporary_file = tempfile.NamedTemporaryFile(
        prefix=".coc-",
        suffix=".xlsx",
        dir=str(output_path.parent),
        delete=False,
    )
    temporary_path = Path(temporary_file.name)
    temporary_file.close()

    workbook = None
    try:
        shutil.copy2(template, temporary_path)
        workbook = load_workbook(temporary_path)
        if "OSX Template" not in workbook.sheetnames:
            raise ValueError("The COC template does not contain an 'OSX Template' sheet.")
        sheet = workbook["OSX Template"]

        # These are the top-left cells of the merged template ranges.
        sheet["B4"] = part_number
        sheet["D4"] = serial
        sheet["F4"] = export_date or date.today()

        records_by_channel = {
            record.channel: record
            for record in measurements
            if 1 <= record.channel <= 48
        }
        for channel, row in enumerate(_measurement_rows(), start=1):
            record = records_by_channel.get(channel)
            sheet["D%d" % row] = record.loss_1310 if record else None
            sheet["E%d" % row] = record.loss_1550 if record else None

        workbook.save(temporary_path)
        workbook.close()
        workbook = None
        os.replace(temporary_path, output_path)
    finally:
        if workbook is not None:
            workbook.close()
        if temporary_path.exists():
            temporary_path.unlink()

    return output_path
