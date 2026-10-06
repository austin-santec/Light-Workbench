import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook

from coc_export import (
    export_coc,
    find_part_number,
    inspect_coc_template_capacity,
    template_filename_for_channel_count,
)
from infrastructure.coc_export import export_coc as infrastructure_export_coc
from run_data import MeasurementRecord


class CocExportTests(unittest.TestCase):
    def test_legacy_module_reexports_infrastructure_exporter(self):
        self.assertIs(export_coc, infrastructure_export_coc)

    def make_template(self, directory, capacity=48):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Sheet1"
        template = workbook.create_sheet("OSX Template")
        template["B1"] = "Template title"
        template.merge_cells("B4:C4")
        template.merge_cells("D4:E4")
        template["B4"] = "old part"
        template["D4"] = "old serial"
        template["F4"] = date(2020, 1, 1)
        rows = list(range(11, 56))
        if capacity == 48:
            rows += list(range(67, 70))
        for channel, row in enumerate(rows, start=1):
            template["B%d" % row] = channel
            template["D%d" % row] = "old 1310"
            template["E%d" % row] = "old 1550"
        path = Path(directory) / "template.xlsx"
        workbook.save(path)
        return path

    def test_exports_partial_readings_to_expected_cells_and_preserves_template(self):
        with tempfile.TemporaryDirectory() as directory:
            template = self.make_template(directory)
            output = export_coc(
                template,
                directory,
                [
                    MeasurementRecord(1, 1.2345, 2.3456),
                    MeasurementRecord(48, 4.5678, 5.6789),
                    MeasurementRecord(49, 9.0, 9.0),
                ],
                "OSX-150-TEST",
                "00123",
                export_date=date(2026, 9, 16),
                tested_by="AJ",
                export_timestamp=datetime(2026, 9, 16, 11, 26, 20),
            )
            workbook = load_workbook(output)
            sheet = workbook["OSX Template"]

            self.assertEqual(output.name, "COC OSX-150 00123_260916-112620.xlsx")
            self.assertEqual(sheet["B4"].value, "OSX-150-TEST")
            self.assertEqual(sheet["D4"].value, "00123")
            self.assertEqual(sheet["F4"].value.date(), date(2026, 9, 16))
            self.assertEqual(sheet["G4"].value, "AJ")
            self.assertEqual(sheet["D11"].value, 1.2345)
            self.assertEqual(sheet["E11"].value, 2.3456)
            self.assertIsNone(sheet["D12"].value)
            self.assertEqual(sheet["D69"].value, 4.5678)
            self.assertEqual(sheet["E69"].value, 5.6789)
            self.assertEqual(sheet["B1"].value, "Template title")

    def test_filename_collision_gets_suffix(self):
        with tempfile.TemporaryDirectory() as directory:
            template = self.make_template(directory)
            timestamp = datetime(2026, 9, 16, 11, 26, 20)
            first = export_coc(
                template,
                directory,
                [],
                "PART",
                "12345",
                export_timestamp=timestamp,
            )
            second = export_coc(
                template,
                directory,
                [],
                "PART",
                "12345",
                export_timestamp=timestamp,
            )
            self.assertEqual(first.name, "COC OSX-150 12345_260916-112620.xlsx")
            self.assertEqual(
                second.name,
                "COC OSX-150 12345_260916-112620 (2).xlsx",
            )

    def test_bundled_templates_are_detected_by_capacity(self):
        template_root = Path(__file__).resolve().parents[1] / "Templates"
        template_45 = template_root / "OSX-150 Single Mode COC Template 2 (45max).xlsx"
        template_48 = template_root / "OSX-150 Single Mode COC Template 1 (48max).xlsx"
        self.assertEqual(inspect_coc_template_capacity(template_45), 45)
        self.assertEqual(inspect_coc_template_capacity(template_48), 48)
        self.assertEqual(template_filename_for_channel_count(45), template_45.name)
        self.assertEqual(template_filename_for_channel_count(46), template_48.name)

    def test_template_selection_boundaries(self):
        template_45 = "OSX-150 Single Mode COC Template 2 (45max).xlsx"
        template_48 = "OSX-150 Single Mode COC Template 1 (48max).xlsx"
        for count in (1, 36, 45):
            with self.subTest(count=count):
                self.assertEqual(template_filename_for_channel_count(count), template_45)
        for count in (46, 47, 48):
            with self.subTest(count=count):
                self.assertEqual(template_filename_for_channel_count(count), template_48)
        for count in (0, 49):
            with self.subTest(count=count):
                with self.assertRaises(ValueError):
                    template_filename_for_channel_count(count)

    def test_export_trims_unused_rows_and_keeps_source_unchanged(self):
        template = (
            Path(__file__).resolve().parents[1]
            / "Templates"
            / "OSX-150 Single Mode COC Template 2 (45max).xlsx"
        )
        original = template.read_bytes()
        source_workbook = load_workbook(template)
        source_sheet = source_workbook["OSX Template"]
        expected_style = source_sheet["B1"].style_id
        expected_width = source_sheet.column_dimensions["B"].width
        expected_orientation = source_sheet.page_setup.orientation
        source_workbook.close()
        with tempfile.TemporaryDirectory() as directory:
            output = export_coc(
                template,
                directory,
                [MeasurementRecord(channel, 1.0, 1.1) for channel in range(1, 37)],
                "PART",
                "12345",
                front_panel_channel_count=36,
            )
            workbook = load_workbook(output)
            sheet = workbook["OSX Template"]
            self.assertEqual(sheet.max_row, 46)
            self.assertEqual(sheet["B46"].value, 36)
            self.assertEqual(sheet.print_area, "'OSX Template'!$B$1:$H$46")
            self.assertIn("F11:F46", {str(item) for item in sheet.merged_cells.ranges})
            self.assertEqual(sheet["B1"].style_id, expected_style)
            self.assertEqual(sheet.column_dimensions["B"].width, expected_width)
            self.assertEqual(sheet.page_setup.orientation, expected_orientation)
            workbook.close()
        self.assertEqual(template.read_bytes(), original)

    def test_all_supported_boundary_counts_trim_to_expected_last_row(self):
        template_root = Path(__file__).resolve().parents[1] / "Templates"
        cases = (
            (1, "OSX-150 Single Mode COC Template 2 (45max).xlsx", 11),
            (36, "OSX-150 Single Mode COC Template 2 (45max).xlsx", 46),
            (45, "OSX-150 Single Mode COC Template 2 (45max).xlsx", 55),
            (46, "OSX-150 Single Mode COC Template 1 (48max).xlsx", 67),
            (47, "OSX-150 Single Mode COC Template 1 (48max).xlsx", 68),
            (48, "OSX-150 Single Mode COC Template 1 (48max).xlsx", 69),
        )
        for count, filename, last_row in cases:
            with self.subTest(count=count), tempfile.TemporaryDirectory() as directory:
                template = template_root / filename
                before = template.read_bytes()
                output = export_coc(
                    template,
                    directory,
                    [
                        MeasurementRecord(channel, channel / 100, channel / 90)
                        for channel in range(1, count + 1)
                    ],
                    "PART",
                    str(10000 + count),
                    front_panel_channel_count=count,
                )
                workbook = load_workbook(output)
                sheet = workbook["OSX Template"]
                self.assertEqual(sheet.max_row, last_row)
                self.assertEqual(sheet["B%d" % last_row].value, count)
                self.assertEqual(sheet["D%d" % last_row].value, count / 100)
                self.assertEqual(
                    sheet.print_area,
                    "'OSX Template'!$B$1:$H$%d" % last_row,
                )
                for merged_range in sheet.merged_cells.ranges:
                    self.assertLessEqual(merged_range.max_row, last_row)
                workbook.close()
                self.assertEqual(template.read_bytes(), before)

    def test_bundled_template_drawings_are_preserved(self):
        template = (
            Path(__file__).resolve().parents[1]
            / "Templates"
            / "OSX-150 Single Mode COC Template 1 (48max).xlsx"
        )
        with tempfile.TemporaryDirectory() as directory:
            output = export_coc(
                template,
                directory,
                [MeasurementRecord(channel, 1.0, 1.1) for channel in range(1, 49)],
                "PART",
                "12345",
                front_panel_channel_count=48,
            )
            with ZipFile(template) as template_zip, ZipFile(output) as output_zip:
                drawing_parts = [
                    name
                    for name in template_zip.namelist()
                    if name.startswith("xl/drawings/") or name.startswith("xl/media/")
                ]
                self.assertTrue(drawing_parts)
                for name in drawing_parts:
                    self.assertIn(name, output_zip.namelist())
                    self.assertEqual(template_zip.read(name), output_zip.read(name))

    def test_malformed_channel_map_is_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            template = self.make_template(directory, capacity=45)
            workbook = load_workbook(template)
            workbook["OSX Template"]["B11"] = 99
            workbook.save(template)
            workbook.close()

            with self.assertRaisesRegex(ValueError, "invalid channel map"):
                inspect_coc_template_capacity(template)

    def test_unusable_output_path_fails_without_changing_inputs(self):
        with tempfile.TemporaryDirectory() as directory:
            template = self.make_template(directory, capacity=45)
            records = [MeasurementRecord(1, 1.0, 1.1)]
            before = tuple(records)
            output_target = Path(directory) / "not-a-directory"
            output_target.write_text("occupied", encoding="utf-8")

            with self.assertRaises(OSError):
                export_coc(
                    template,
                    output_target,
                    records,
                    "PART",
                    "12345",
                    front_panel_channel_count=1,
                )
            self.assertEqual(tuple(records), before)

    def test_45_capacity_template_rejects_46_channels(self):
        with tempfile.TemporaryDirectory() as directory:
            template = self.make_template(directory, capacity=45)
            with self.assertRaisesRegex(ValueError, "supports only 45"):
                export_coc(
                    template,
                    directory,
                    [MeasurementRecord(channel, 1.0, 1.1) for channel in range(1, 47)],
                    "PART",
                    "12345",
                    front_panel_channel_count=46,
                )

    def test_finds_part_number_from_serial_folder(self):
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory) / "SN00123_OSX-150-1A-008-09-FA-00B-1H"
            folder.mkdir()
            self.assertEqual(find_part_number("00123", directory), "OSX-150-1A-008-09-FA-00B-1H")
            self.assertEqual(find_part_number("SN00123", directory), "OSX-150-1A-008-09-FA-00B-1H")


if __name__ == "__main__":
    unittest.main()
